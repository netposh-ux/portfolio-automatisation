"""Read PDF invoices and put their data into one Excel workbook.

For each invoice: number, date, supplier, subtotal, taxes and total, plus every line item.
Each invoice is also checked: if the lines plus taxes don't add up to the total, it's flagged.

Example: python extract.py --folder invoices --output invoices.xlsx
"""
import argparse
import re
from datetime import datetime
from pathlib import Path

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

HERE = Path(__file__).resolve().parent
MONEY_FORMAT = '"$"#,##0.00'
FLAG_FILL = PatternFill("solid", fgColor="F8CBAD")

NUMBER_PATTERNS = [r"Invoice #:\s*(\S+)", r"Invoice No\.\s*(\S+)"]
DATE_PATTERNS = [r"Date:\s*(.+)", r"Issued on\s*(.+)"]
DATE_FORMATS = ["%Y-%m-%d", "%B %d, %Y", "%d/%m/%Y"]
ITEM_LINE = re.compile(r"^(.+?)\s+(\d+)\s+\$([\d,]+\.\d{2})\s+\$([\d,]+\.\d{2})$")


def amount(text):
    return float(text.replace(",", ""))


def first_match(patterns, text):
    for pattern in patterns:
        found = re.search(pattern, text)
        if found:
            return found.group(1).strip()
    return None


def parse_date(text):
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except (TypeError, ValueError):
            continue
    return None


def read_invoice(path):
    with pdfplumber.open(path) as pdf:
        text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    lines = [line.strip() for line in text.splitlines() if line.strip()]

    def total_for(label):
        found = re.search(rf"^{re.escape(label)}\S*.*?\$([\d,]+\.\d{{2}})$", text, re.M)
        return amount(found.group(1)) if found else None

    items = []
    for line in lines:
        match = ITEM_LINE.match(line)
        if match:
            description, qty, unit, line_amount = match.groups()
            items.append({"description": description, "qty": int(qty),
                          "unit_price": amount(unit), "amount": amount(line_amount)})

    invoice = {
        "file": path.name,
        "number": first_match(NUMBER_PATTERNS, text),
        "date": parse_date(first_match(DATE_PATTERNS, text)),
        "supplier": lines[0] if lines else None,
        "subtotal": total_for("Subtotal"),
        "taxes": total_for("Taxes"),
        "total": total_for("Total"),
        "items": items,
    }
    expected = round(sum(i["amount"] for i in items) + (invoice["taxes"] or 0), 2)
    invoice["check"] = "OK" if invoice["total"] == expected else f"Lines + taxes = ${expected:,.2f}"
    return invoice


def style_header(sheet, widths):
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for column, width in zip("ABCDEFGHIJ", widths):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def save_excel(invoices, path):
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Invoices"
    summary.append(["Number", "Date", "Supplier", "Subtotal", "Taxes", "Total", "Check", "File"])
    for inv in sorted(invoices, key=lambda i: (i["date"] or datetime.min.date())):
        summary.append([inv["number"], inv["date"], inv["supplier"], inv["subtotal"],
                        inv["taxes"], inv["total"], inv["check"], inv["file"]])
        if inv["check"] != "OK":
            for cell in summary[summary.max_row]:
                cell.fill = FLAG_FILL
    for row in summary.iter_rows(min_row=2):
        row[1].number_format = "yyyy-mm-dd"
        for cell in row[3:6]:
            cell.number_format = MONEY_FORMAT
    style_header(summary, (12, 12, 28, 12, 12, 12, 26, 24))

    lines = workbook.create_sheet("Line items")
    lines.append(["Invoice", "Supplier", "Description", "Qty", "Unit price", "Amount"])
    for inv in invoices:
        for item in inv["items"]:
            lines.append([inv["number"], inv["supplier"], item["description"], item["qty"],
                          item["unit_price"], item["amount"]])
    for row in lines.iter_rows(min_row=2, min_col=5, max_col=6):
        for cell in row:
            cell.number_format = MONEY_FORMAT
    style_header(lines, (12, 28, 36, 6, 12, 12))
    workbook.save(path)


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--folder", type=Path, default=HERE / "invoices")
    parser.add_argument("--output", default="invoices.xlsx")
    args = parser.parse_args()

    files = sorted(args.folder.glob("*.pdf"))
    if not files:
        raise SystemExit(f"No PDF files in {args.folder} (run make_sample_invoices.py first)")

    invoices = []
    for path in files:
        inv = read_invoice(path)
        invoices.append(inv)
        print(f"{inv['number']:<10} {inv['date']}  {inv['supplier']:<26} ${inv['total']:>9,.2f}  {inv['check']}")

    save_excel(invoices, args.output)
    flagged = sum(i["check"] != "OK" for i in invoices)
    print(f"\n{len(invoices)} invoices, {sum(len(i['items']) for i in invoices)} line items, "
          f"{flagged} to double-check -> {args.output}")


if __name__ == "__main__":
    main()
