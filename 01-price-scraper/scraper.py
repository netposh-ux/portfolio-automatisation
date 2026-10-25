"""Scrape products from an online catalog and save them to an Excel file.

Example: python scraper.py --pages 3 --output prices.xlsx
"""
import argparse
import time
from datetime import datetime
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

START_URL = "https://books.toscrape.com/catalogue/page-1.html"
RATINGS = {"One": 1, "Two": 2, "Three": 3, "Four": 4, "Five": 5}
PAUSE_SECONDS = 1.0  # be polite: at most one page per second


def read_page(session, url):
    """Return (products, next_page_url) for one catalog page."""
    response = session.get(url, timeout=15)
    response.raise_for_status()
    response.encoding = "utf-8"
    soup = BeautifulSoup(response.text, "html.parser")

    products = []
    for article in soup.select("article.product_pod"):
        link = article.h3.a
        price_text = article.select_one(".price_color").get_text(strip=True)
        rating_classes = article.select_one(".star-rating")["class"]
        products.append({
            "title": link["title"],
            "price": float(price_text.lstrip("£")),
            "in_stock": "In stock" in article.select_one(".availability").get_text(),
            "rating": next((RATINGS[c] for c in rating_classes if c in RATINGS), None),
            "url": urljoin(url, link["href"]),
        })

    next_link = soup.select_one("li.next a")
    return products, (urljoin(url, next_link["href"]) if next_link else None)


def scrape(page_count):
    """Walk up to page_count pages and return every product found."""
    session = requests.Session()
    session.headers["User-Agent"] = "portfolio-automation/1.0 (demo)"
    products, url = [], START_URL
    for number in range(1, page_count + 1):
        if not url:
            break
        page, url = read_page(session, url)
        products.extend(page)
        print(f"Page {number}: {len(page)} products")
        if url and number < page_count:
            time.sleep(PAUSE_SECONDS)
    return products


def save_excel(products, path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Products"
    sheet.append(["Title", "Price (£)", "In stock", "Rating /5", "Link"])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")

    for p in sorted(products, key=lambda p: p["price"]):
        sheet.append([p["title"], p["price"], "Yes" if p["in_stock"] else "No", p["rating"], p["url"]])

    for column, width in zip("ABCDE", (60, 10, 10, 10, 70)):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    workbook.save(path)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--pages", type=int, default=2, help="number of pages to scrape (default: 2)")
    parser.add_argument("--output", default=f"prices_{datetime.now():%Y-%m-%d}.xlsx")
    args = parser.parse_args()

    products = scrape(args.pages)
    save_excel(products, args.output)
    average = sum(p["price"] for p in products) / len(products)
    print(f"Saved {len(products)} products to {args.output} (average price: £{average:.2f})")


if __name__ == "__main__":
    main()
