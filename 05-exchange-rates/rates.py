"""Pull daily exchange rates from a public API, keep them in a SQLite database,
and export an Excel file and a chart.

Only the missing days are downloaded, so you can run it every morning from cron
and it just adds the new day.

Example: python rates.py --base CAD --currencies USD EUR GBP --days 90
"""
import argparse
import sqlite3
from datetime import date, timedelta
from pathlib import Path

import matplotlib
import requests
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font, PatternFill

matplotlib.use("Agg")  # no screen needed
import matplotlib.pyplot as plt  # noqa: E402

HERE = Path(__file__).resolve().parent
API_URL = "https://api.frankfurter.dev/v1/{start}..{end}"  # European Central Bank data, no API key


def open_database(path):
    db = sqlite3.connect(path)
    db.execute("""CREATE TABLE IF NOT EXISTS rates (
                      day TEXT, base TEXT, currency TEXT, rate REAL,
                      PRIMARY KEY (day, base, currency))""")
    return db


def latest_day(db, base):
    row = db.execute("SELECT MAX(day) FROM rates WHERE base = ?", (base,)).fetchone()
    return date.fromisoformat(row[0]) if row[0] else None


def download(base, currencies, start, end):
    response = requests.get(API_URL.format(start=start, end=end),
                            params={"base": base, "symbols": ",".join(currencies)}, timeout=20)
    response.raise_for_status()
    return response.json()["rates"]  # {"2026-09-01": {"USD": 0.72, ...}, ...}


def update(db, base, currencies, days):
    """Download only what the database doesn't have yet. Returns the number of new rows."""
    start = date.today() - timedelta(days=days)
    last = latest_day(db, base)
    if last and last >= start:
        start = last + timedelta(days=1)
    if start > date.today():
        return 0
    # On weekends the API answers with Friday's rates, which we already have: skip them
    rows = [(day, base, currency, rate)
            for day, values in download(base, currencies, start, date.today()).items()
            if day >= start.isoformat()
            for currency, rate in values.items()]
    db.executemany("INSERT OR REPLACE INTO rates VALUES (?, ?, ?, ?)", rows)
    db.commit()
    return len(rows)


def load(db, base, currencies, days):
    since = (date.today() - timedelta(days=days)).isoformat()
    table = {}
    for day, currency, rate in db.execute(
            "SELECT day, currency, rate FROM rates WHERE base = ? AND day >= ? ORDER BY day", (base, since)):
        if currency in currencies:
            table.setdefault(day, {})[currency] = rate
    return table


def save_excel(table, base, currencies, path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = f"{base} rates"
    sheet.append(["Date"] + [f"1 {base} in {c}" for c in currencies])
    for day, values in table.items():
        sheet.append([date.fromisoformat(day)] + [values.get(c) for c in currencies])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for row in sheet.iter_rows(min_row=2):
        row[0].number_format = "yyyy-mm-dd"
        for cell in row[1:]:
            cell.number_format = "0.0000"
    sheet.column_dimensions["A"].width = 12
    sheet.freeze_panes = "A2"

    chart = LineChart()
    chart.title = f"Value of 1 {base}"
    chart.height, chart.width = 9, 22
    chart.add_data(Reference(sheet, min_col=2, max_col=len(currencies) + 1, min_row=1, max_row=sheet.max_row),
                   titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=sheet.max_row))
    sheet.add_chart(chart, f"{chr(ord('B') + len(currencies) + 1)}2")
    workbook.save(path)


def save_chart(table, base, currencies, path):
    days = [date.fromisoformat(d) for d in table]
    fig, ax = plt.subplots(figsize=(9, 4.2), dpi=110)
    for currency in currencies:
        ax.plot(days, [table[d].get(currency) for d in table], label=currency, linewidth=1.8)
    ax.set_title(f"Value of 1 {base}, last {len(days)} business days")
    ax.grid(alpha=0.3)
    ax.legend(frameon=False, loc="center left", bbox_to_anchor=(1, 0.5))
    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    fig.autofmt_xdate()
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--base", default="CAD")
    parser.add_argument("--currencies", nargs="+", default=["USD", "EUR", "GBP"])
    parser.add_argument("--days", type=int, default=90)
    parser.add_argument("--database", type=Path, default=HERE / "rates.db")
    parser.add_argument("--output", default="rates.xlsx")
    parser.add_argument("--chart", default="rates.png")
    args = parser.parse_args()

    db = open_database(args.database)
    print(f"Downloaded {update(db, args.base, args.currencies, args.days)} new rate(s)")
    table = load(db, args.base, args.currencies, args.days)
    if not table:
        raise SystemExit("No data for that period")

    first, last = table[min(table)], table[max(table)]
    print(f"\n1 {args.base} on {max(table)}:")
    for c in args.currencies:
        change = (last[c] / first[c] - 1) * 100
        print(f"  {c}  {last[c]:.4f}   ({change:+.2f}% since {min(table)})")

    save_excel(table, args.base, args.currencies, args.output)
    save_chart(table, args.base, args.currencies, args.chart)
    print(f"\nSaved {args.output} and {args.chart}")


if __name__ == "__main__":
    main()
