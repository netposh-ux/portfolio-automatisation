"""Merge every CSV in a folder into one Excel report with totals and charts.

Example: python merge.py --folder data --output report.xlsx
"""
import argparse
import csv
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill

HERE = Path(__file__).resolve().parent
MONEY_FORMAT = '"$"#,##0.00'


def number(text):
    """Accept both 4.50 and 4,50."""
    return float(text.strip().replace(",", "."))


def read_csv(path):
    """Read a CSV, detecting the separator (comma or semicolon)."""
    with open(path, newline="", encoding="utf-8-sig") as f:
        sample = f.read(2048)
        f.seek(0)
        dialect = csv.Sniffer().sniff(sample, delimiters=",;")
        rows = []
        for line_number, row in enumerate(csv.DictReader(f, dialect=dialect), start=2):
            try:
                quantity = int(row["quantity"])
                price = number(row["unit_price"])
            except (KeyError, ValueError, AttributeError):
                print(f"  skipped line {line_number} of {path.name} (invalid value)")
                continue
            rows.append({
                "date": row["date"].strip(),
                "region": row["region"].strip(),
                "product": row["product"].strip(),
                "quantity": quantity,
                "unit_price": price,
                "total": round(quantity * price, 2),
            })
        return rows


def style_header(sheet):
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    sheet.freeze_panes = "A2"


def totals_sheet(workbook, title, key, label, rows):
    sums = defaultdict(lambda: [0, 0.0])
    for r in rows:
        sums[r[key]][0] += r["quantity"]
        sums[r[key]][1] += r["total"]

    sheet = workbook.create_sheet(title)
    sheet.append([label, "Quantity", "Sales"])
    for name, (quantity, sales) in sorted(sums.items(), key=lambda s: -s[1][1]):
        sheet.append([name, quantity, round(sales, 2)])
    last = sheet.max_row
    sheet.append(["TOTAL", f"=SUM(B2:B{last})", f"=SUM(C2:C{last})"])
    sheet[f"A{last + 1}"].font = Font(bold=True)
    for row in sheet.iter_rows(min_row=2, min_col=3, max_col=3):
        row[0].number_format = MONEY_FORMAT
    sheet.column_dimensions["A"].width = 16
    sheet.column_dimensions["C"].width = 14
    style_header(sheet)

    chart = BarChart()
    chart.title = f"Sales by {label.lower()}"
    chart.legend = None
    chart.add_data(Reference(sheet, min_col=3, min_row=1, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=last))
    sheet.add_chart(chart, "E2")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--folder", type=Path, default=HERE / "data")
    parser.add_argument("--output", default="sales_report.xlsx")
    args = parser.parse_args()

    files = sorted(args.folder.glob("*.csv"))
    if not files:
        raise SystemExit(f"No CSV files in {args.folder} (run make_sample_data.py first)")

    rows = []
    for path in files:
        content = read_csv(path)
        print(f"{path.name}: {len(content)} rows")
        rows.extend(content)

    workbook = Workbook()
    detail = workbook.active
    detail.title = "Detail"
    columns = ["date", "region", "product", "quantity", "unit_price", "total"]
    detail.append(columns)
    for r in sorted(rows, key=lambda r: (r["date"], r["region"], r["product"])):
        detail.append([r[c] for c in columns])
    for row in detail.iter_rows(min_row=2, min_col=5, max_col=6):
        for cell in row:
            cell.number_format = MONEY_FORMAT
    detail.auto_filter.ref = detail.dimensions
    style_header(detail)

    totals_sheet(workbook, "By region", "region", "Region", rows)
    totals_sheet(workbook, "By product", "product", "Product", rows)
    workbook.save(args.output)

    total = sum(r["total"] for r in rows)
    print(f"Merged {len(rows)} rows, total sales ${total:,.2f} -> {args.output}")


if __name__ == "__main__":
    main()
